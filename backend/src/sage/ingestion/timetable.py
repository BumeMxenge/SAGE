# Reads the department's timetable workbook, by CP0's readings. Each sheet is a grid of year blocks,
# with the days of each semester across and the handbook's periods down. A block of merged cells is
# one meeting, from its first period's start to its last period's end, and a cell naming several
# sections, as in "(2278 or 2483)", is parallel groups at one time, so one meeting. The reader is
# strict: every cell on a sheet it reads is placed, or the run stops and names it.
import re
from collections.abc import Mapping, Sequence
from datetime import time
from pathlib import Path
from string import ascii_uppercase
from typing import Literal
from uuid import UUID

import openpyxl
from openpyxl.cell.cell import Cell
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.worksheet import Worksheet

from sage.domain.base import Row
from sage.domain.rules import (
    TERM_BY_SUFFIX,
    TermCode,
    TimetableSlot,
    TimetableSlotCurriculum,
)
from sage.ingestion.sources import Alternatives, SourceError, row_id

Semester = Literal["sem1", "sem2"]

# The EBE handbook's key (p. 10): each period lasts 45 minutes, and the meridian 13:00 to 14:00
PERIODS = {
    "1": (time(8), time(8, 45)),
    "2": (time(9), time(9, 45)),
    "3": (time(10), time(10, 45)),
    "4": (time(11), time(11, 45)),
    "5": (time(12), time(12, 45)),
    "M": (time(13), time(14)),
    "6": (time(14), time(14, 45)),
    "7": (time(15), time(15, 45)),
    "8": (time(16), time(16, 45)),
    "9": (time(17), time(17, 45)),
}
DAYS = {"MONDAY": 1, "TUESDAY": 2, "WEDNESDAY": 3, "THURSDAY": 4, "FRIDAY": 5}
SEMESTERS: dict[str, Semester] = {"Semester 1": "sem1", "Semester 2": "sem2"}
# The terms a semester's columns can hold: the semester and its quarters
WITHIN: dict[Semester, tuple[TermCode, ...]] = {
    "sem1": ("sem1", "q1", "q2"),
    "sem2": ("sem2", "q3", "q4"),
}
YEAR = re.compile(r"(\d)(?:ST|ND|RD|TH) YEAR")
# A meeting: its code, with "EEE3097/8S" meaning EEE3097S and EEE3098S, then any sections
MEETING = re.compile(r"([A-Z]{3}\d{3})(\d)(?:/(\d))?([A-Z])\s*(?:\([^)]*\))?")


class Printed(Row):
    # A meeting as one sheet prints it
    sheet: str
    cells: str  # such as "C7:C8"
    codes: tuple[str, ...]
    year: int | None  # the block's year of study; None on a sheet of labs
    semester: Semester
    day: int
    first: str  # the first period, as the sheet labels it
    last: str


class Note(Row):
    sheet: str
    cell: str
    text: str


class Workbook(Row):
    printed: tuple[Printed, ...]
    notes: tuple[Note, ...]


class Timetable(Row):
    slots: tuple[TimetableSlot, ...]
    links: tuple[TimetableSlotCurriculum, ...]


def read(path: Path, sheets: Mapping[str, Sequence[str]]) -> Workbook:
    book = openpyxl.load_workbook(path, data_only=True)
    printed: list[Printed] = []
    notes: list[Note] = []
    for name, curricula in sheets.items():
        sheet = book[name]
        if sheet.sheet_state != "visible":
            raise SourceError(
                f"The {name} sheet is hidden, and a hidden sheet is never read"
            )
        _sheet(sheet, not curricula, printed, notes)
    return Workbook(printed=tuple(printed), notes=tuple(notes))


def _sheet(
    sheet: Worksheet, labs: bool, printed: list[Printed], notes: list[Note]
) -> None:
    # The grid's structure first, then its meetings, as a block's periods are only known once all
    # its rows are read. A sheet of labs titles its blocks by lab, not year of study.
    merged = {cells.coord.split(":")[0]: cells for cells in sheet.merged_cells.ranges}
    # Each block's first row and year of study
    blocks: list[tuple[int, int | None]] = []
    semesters: dict[int, Semester] = {}  # each semester's time column
    days: dict[int, tuple[Semester, int]] = {}
    periods: dict[int, str] = {}  # each period row's label
    meetings: list[Cell] = []
    for row in sheet.iter_rows():
        for cell in row:
            value, where = cell.value, f"{sheet.title}!{cell.coordinate}"
            text = value.strip() if isinstance(value, str) else None
            spans = merged.get(cell.coordinate)
            if value is None or text == "":
                continue  # empty, or only spaces, which print nothing
            if text and text.startswith("Note:"):
                notes.append(Note(sheet=sheet.title, cell=cell.coordinate, text=text))
            elif text in SEMESTERS:
                semesters[cell.column] = SEMESTERS[text]
            elif (
                text
                and text.endswith(tuple(SEMESTERS))
                or text in ("Start Time", "Period")
            ):
                pass  # a block's heading over each semester, and the period key's columns
            elif text in DAYS:
                semester = semesters[max(c for c in semesters if c < cell.column)]
                days[cell.column] = (semester, DAYS[text])
            elif cell.column == 1 and spans and spans.max_col > spans.min_col + 7:
                found = YEAR.fullmatch(text or "")
                if not (labs or found):
                    raise SourceError(f"{where} isn't a year of study: {value!r}")
                blocks.append((cell.row, None if labs else int(found[1])))
            elif cell.column in semesters and isinstance(value, time):
                pass  # a period's start, checked against its label beside it
            elif cell.column - 1 in semesters and str(value) in PERIODS:
                start = sheet.cell(cell.row, cell.column - 1).value
                if start != PERIODS[str(value)][0]:
                    raise SourceError(
                        f"{where}: period {value} doesn't start at {start}"
                    )
                periods[cell.row] = str(value)
            elif cell.column in days and text:
                meetings.append(cell)
            else:
                raise SourceError(f"{where} fits no reading: {value!r}")
    for cell in meetings:
        year = next(year for first, year in reversed(blocks) if first < cell.row)
        printed.append(
            _meeting(
                sheet.title,
                cell,
                merged.get(cell.coordinate),
                cell.value.strip(),
                year,
                days,
                periods,
            )
        )


def _meeting(
    sheet: str,
    cell: Cell,
    spans: CellRange | None,
    text: str,
    year: int | None,
    days: dict[int, tuple[Semester, int]],
    periods: dict[int, str],
) -> Printed:
    cells = spans.coord if spans else cell.coordinate
    found = MEETING.fullmatch(text)
    if not found:
        raise SourceError(f"{sheet}!{cells} isn't a meeting: {text!r}")
    stem, digit, other, suffix = found.groups()
    last = spans.max_row if spans else cell.row
    if (spans and spans.max_col != spans.min_col) or not all(
        row in periods for row in range(cell.row, last + 1)
    ):
        raise SourceError(f"{sheet}!{cells} isn't one day's periods")
    semester, day = days[cell.column]
    return Printed(
        sheet=sheet,
        cells=cells,
        codes=tuple(stem + d + suffix for d in (digit, other) if d),
        year=year,
        semester=semester,
        day=day,
        first=periods[cell.row],
        last=periods[last],
    )


def stored(
    workbook: Workbook,
    sheets: Mapping[str, Sequence[str]],
    taking: Mapping[str, set[str]],
    alternatives: Sequence[Alternatives],
    year: int,
    known: set[str],
) -> Timetable:
    # The meetings to store: each printed once, however many sheets print it. A meeting the sheet of
    # every curriculum taking the course prints needs no curriculum links; any other is linked to the
    # curricula whose sheets print it. A sheet of labs names no curricula, so a lab it alone prints
    # is for everyone on the course, and every meeting it prints is a lab.
    found: dict[tuple[str, TermCode, int, time, time], tuple[set[str], list[str]]] = {}
    for meeting in workbook.printed:
        start, end = PERIODS[meeting.first][0], PERIODS[meeting.last][1]
        for code in meeting.codes:
            if code in known:
                key = (code, _term(code, meeting.semester), meeting.day, start, end)
                curricula, places = found.setdefault(key, (set(), []))
                curricula.update(sheets[meeting.sheet])
                places.append(f"{meeting.sheet}!{meeting.cells}")
    choices = {}
    for alternative in alternatives:
        for number, place in enumerate(alternative.cells):
            key = next(
                (
                    k
                    for k, (_, places) in found.items()
                    if k[0] == alternative.course and place in places
                ),
                None,
            )
            if key is None:
                raise SourceError(f"{place} prints no meeting of {alternative.course}")
            choices[key] = (alternative.group, ascii_uppercase[number])
    slots, links = [], []
    for key, (curricula, places) in found.items():
        code, term, day, start, end = key
        group, choice = choices.get(key, (None, None))
        row = {
            "handbook_year": year,
            "course_code": code,
            "term": term,
            "day": day,
            "start_time": start,
            "end_time": end,
            "kind": "lab" if any(not sheets[p.split("!")[0]] for p in places) else None,
            "choice_group": group,
            "choice": choice,
        }
        slot = TimetableSlot(id=row_id(row), **row)
        slots.append(slot)
        if curricula and not curricula >= taking.get(code, set()):
            links += [
                TimetableSlotCurriculum(
                    slot_id=slot.id, handbook_year=year, curriculum_code=c
                )
                for c in sorted(curricula)
            ]
    return Timetable(slots=tuple(slots), links=tuple(links))


def round_trip(
    workbook: Workbook,
    timetable: Timetable,
    sheets: Mapping[str, Sequence[str]],
    years: Mapping[tuple[str, str], int],
    known: set[str],
) -> tuple[str, ...]:
    # Rebuilds each programme sheet from the stored meetings, by the year of study each curriculum
    # takes a course in, and compares it with the sheet less section numbers and notes. A lab only
    # the labs sheet prints is expected on the sheets of every curriculum taking its course. Stored
    # meetings don't say which lab, so the labs sheet is compared as a set of meetings.
    linked: dict[UUID, set[str]] = {}
    for link in timetable.links:
        linked.setdefault(link.slot_id, set()).add(link.curriculum_code)
    cells = [
        (meeting, code)
        for meeting in workbook.printed
        for code in meeting.codes
        if code in known
    ]
    programmes = {_place(m, code) for m, code in cells if sheets[m.sheet]}
    labs = [(m, code) for m, code in cells if not sheets[m.sheet]]
    differences = []
    for sheet, curricula in sheets.items():
        for curriculum in curricula:
            printed = {
                (m.year, *_place(m, code)) for m, code in cells if m.sheet == sheet
            }
            printed |= {
                (years[curriculum, code], *_place(m, code))
                for m, code in labs
                if (curriculum, code) in years and _place(m, code) not in programmes
            }
            rebuilt = {
                (years[curriculum, slot.course_code], *_stored(slot))
                for slot in timetable.slots
                if (curriculum, slot.course_code) in years
                and curriculum in linked.get(slot.id, {curriculum})
            }
            differences += [
                f"{sheet} prints {c}, which no stored meeting gives"
                for c in sorted(printed - rebuilt, key=str)
            ]
            differences += [
                f"{sheet} lacks {c}, which a stored meeting gives"
                for c in sorted(rebuilt - printed, key=str)
            ]
    meetings = {_stored(slot) for slot in timetable.slots}
    differences += [
        f"{m.sheet}!{m.cells}'s {code} is no stored meeting"
        for m, code in labs
        if _place(m, code) not in meetings
    ]
    return tuple(differences)


def _term(code: str, semester: Semester) -> TermCode:
    # A quarter keeps its own term within its semester's columns. Any other code is stored under the
    # semester printed, and the term check compares that with its suffix.
    term = TERM_BY_SUFFIX.get(code[-1])
    return term if term in WITHIN[semester] else semester


def _place(meeting: Printed, code: str) -> tuple[Semester, int, str, str, str]:
    return (meeting.semester, meeting.day, meeting.first, meeting.last, code)


def _stored(
    slot: TimetableSlot,
) -> tuple[Semester, int | None, str | None, str | None, str]:
    # Where a stored meeting would print: its semester, day, and first and last periods
    semester = next(s for s, terms in WITHIN.items() if slot.term in terms)
    first = next(
        (p for p, (start, _) in PERIODS.items() if start == slot.start_time), None
    )
    last = next((p for p, (_, end) in PERIODS.items() if end == slot.end_time), None)
    return (semester, slot.day, first, last, slot.course_code)
