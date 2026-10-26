# What the ingestion tests share: the handbook's typed pages and a PDF set from them, and a
# timetable block and a complementary workbook in the real layouts
from collections.abc import Callable
from datetime import time
from pathlib import Path

import openpyxl
import pytest
from fpdf import FPDF

from sage.ingestion.complementary import MEETING_COLUMNS, OFFERING_COLUMNS
from sage.ingestion.handbook import Handbook, Page
from tests.pdf import split_pages


@pytest.fixture
def handbook_pages() -> list[list[str]]:
    # Seven pages of the 2026 handbook as printed: the start of the Electrical Engineering section and
    # four course entries. Each opens with its running header, a blank line marks each break, and long
    # passages are shortened. A fresh copy each time, so a test can edit it.
    text = (Path(__file__).parents[2] / "fixtures" / "handbook.txt").read_text(
        encoding="utf-8"
    )
    return split_pages(text)


@pytest.fixture
def handbook(handbook_pages) -> Handbook:
    # The typed pages as the PDF layer reads them: the number each running header ends with, and the
    # body below it
    return Handbook(
        [
            Page(number=int(header.split()[-1]), lines=tuple(body))
            for header, *body in handbook_pages
        ]
    )


@pytest.fixture(scope="session")
def handbook_pdf() -> Callable[[Path, list[list]], Path]:
    # Sets typed pages as the handbook does: each page's first line as the running header at the top,
    # the rest from 60 points down in 8-point type, with a blank line as a gap. A (text, size) pair is
    # set in that size. A4, so the longest typed line fits.
    def make(path: Path, typed: list[list]) -> Path:
        pdf = FPDF(unit="pt", format="A4")
        pdf.set_auto_page_break(False)
        for header, *body in typed:
            pdf.add_page()
            pdf.set_font("Helvetica", size=8)
            if header:
                pdf.set_xy(36, 28)
                pdf.cell(text=header)
            pdf.set_y(60)
            for line in body:
                text, size = line if isinstance(line, tuple) else (line, 8)
                pdf.set_font("Helvetica", size=size)
                if text:
                    pdf.cell(h=size + 1, text=text, new_x="LMARGIN", new_y="NEXT")
                else:
                    pdf.ln(9)
        pdf.output(str(path))
        return path

    return make


@pytest.fixture(scope="session")
def timetable_block() -> Callable:
    # Lays out a block at the top of a sheet as the 2026 timetable workbook does: its title across
    # the grid, each semester's heading and days, and the handbook's period key (p. 10) down each
    # semester's side, with periods 1 to 5, the meridian and 6 to 9 from row 5
    key = (("1", 8), ("2", 9), ("3", 10), ("4", 11), ("5", 12), ("M", 13), ("6", 14))
    key += (("7", 15), ("8", 16), ("9", 17))
    days = ("MONDAY", "TUESDAY", "WEDNESDAY", "THURSDAY", "FRIDAY")

    def block(sheet, title: str) -> None:
        sheet["A1"] = title
        sheet.merge_cells("A1:O1")
        for column, semester in ((1, "Semester 1"), (9, "Semester 2")):
            sheet.cell(2, column, f"Mainstream EE - {semester}")
            sheet.cell(3, column, semester)
            for offset, day in enumerate(days):
                sheet.cell(3, column + 2 + offset, day)
            sheet.cell(4, column, "Start Time")
            sheet.cell(4, column + 1, "Period")
            for row, (period, hour) in enumerate(key, start=5):
                sheet.cell(row, column, time(hour))
                sheet.cell(row, column + 1, period if period == "M" else int(period))

    return block


@pytest.fixture(scope="session")
def complementary_xlsx() -> Callable:
    # Builds a complementary workbook in the real one's layout: each sheet's title, its header, then
    # its rows, with one cell changed afterwards where a test plants a fault
    def make(path: Path, rows: dict[str, list[tuple]], change=None) -> Path:
        book = openpyxl.Workbook()
        book.remove(book.active)
        for name, sheet_rows in rows.items():
            sheet = book.create_sheet(name)
            sheet.append([f"The {name} sheet's title"])
            sheet.append(
                list(MEETING_COLUMNS if name == "Meetings" else OFFERING_COLUMNS)
            )
            for row in sheet_rows:
                sheet.append(list(row))
        if change:
            name, cell, value = change
            book[name][cell] = value
        book.save(path)
        return path

    return make
