# Shared pytest fixtures for every backend test
import json
import os
import tempfile
from collections.abc import Callable
from datetime import time
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient
from fpdf import FPDF

from sage.domain.rules import ResultSymbol, RuleSet
from sage.ingestion.complementary import MEETING_COLUMNS, OFFERING_COLUMNS
from sage.ingestion.handbook import Handbook, Page
from sage.main import create_app
from sage.services.ingest import approved_snapshot
from sage.services.seed import SEED_FILE


@pytest.hookimpl(tryfirst=True)
def pytest_configure(config: pytest.Config) -> None:
    # make mutation runs many test sessions at once, so each takes a temporary folder of its own in
    # SAGE_TEST_TEMP. Their default folders share one pytest-current link, which every session
    # remakes as it ends, and a session that crashes there would count as a caught mutant.
    root = os.environ.get("SAGE_TEST_TEMP")
    if root and config.option.basetemp is None:
        config.option.basetemp = tempfile.mkdtemp(dir=root)


@pytest.fixture
def client() -> TestClient:
    return TestClient(create_app())


@pytest.fixture(scope="session")
def seed_file() -> Path:
    return SEED_FILE


@pytest.fixture(scope="session")
def result_symbols() -> tuple[ResultSymbol, ...]:
    # UCT's 26 symbols, as the rule-tables migration stores them, for the in-memory rulebook. A copy,
    # as the unit tests have no database, and test_the_symbols_are_ucts_26 compares it with the
    # migration's rows on every make check-db, so the two can't drift apart.
    path = Path(__file__).parent / "fixtures" / "result_symbols.json"
    rows = json.loads(path.read_text(encoding="utf-8"))
    return tuple(ResultSymbol.model_validate(row) for row in rows)


@pytest.fixture(scope="session")
def approved_2026() -> RuleSet:
    # The approved 2026 rulebook, data/rules/2026/rulebook.json, read once. A test that needs rules
    # loads it into a store of its own, so no test can change another's.
    return approved_snapshot(2026).rules


@pytest.fixture(scope="session")
def fake_account(seed_file) -> dict[str, str]:
    # Mbali's made-up email, which the seed file keeps for the tests and sage sign-in
    return json.loads(seed_file.read_text(encoding="utf-8"))["fake_account"]


@pytest.fixture(scope="session")
def mbali_counts() -> dict[str, int]:
    # The rows per table the seed file holds, written out by hand rather than counted from the file
    return {
        "students": 1,
        "course_history": 34,
        "favourites": 6,
        "plan_sets": 1,
        "plans": 10,
        "plan_courses": 80,
        "plan_reasons": 6,
        "chats": 2,
        "messages": 4,
        "message_citations": 0,
    }


@pytest.fixture
def transcript_pages() -> list[list[str]]:
    # The made-up transcript in the real layout, page by page, as pdf_pages reads one: Mbali's
    # identity and her seeded lines. A fresh copy each time, so a test can edit it.
    text = (Path(__file__).parent / "fixtures" / "transcript.txt").read_text(
        encoding="utf-8"
    )
    return [page.strip("\n").split("\n") for page in text.split("\f")]


@pytest.fixture
def handbook_pages() -> list[list[str]]:
    # Seven pages of the 2026 handbook as printed: the start of the Electrical Engineering section and
    # four course entries. Each opens with its running header, a blank line marks each break, and long
    # passages are shortened. A fresh copy each time, so a test can edit it.
    text = (Path(__file__).parent / "fixtures" / "handbook.txt").read_text(
        encoding="utf-8"
    )
    return [page.strip("\n").split("\n") for page in text.split("\f")]


@pytest.fixture
def handbook(handbook_pages) -> Handbook:
    # The typed pages as the PDF layer reads them: the number each running header ends with, and the
    # body below it
    return Handbook(
        [
            Page(number=int(header.split()[-1]), lines=tuple(body))
            for header, *body in handbook_pages
        ]
    )


@pytest.fixture(scope="session")
def handbook_pdf() -> Callable[[Path, list[list]], Path]:
    # Sets typed pages as the handbook does: each page's first line as the running header at the top,
    # the rest from 60 points down in 8-point type, with a blank line as a gap. A (text, size) pair is
    # set in that size. A4, so the longest typed line fits.
    def make(path: Path, typed: list[list]) -> Path:
        pdf = FPDF(unit="pt", format="A4")
        pdf.set_auto_page_break(False)
        for header, *body in typed:
            pdf.add_page()
            pdf.set_font("Helvetica", size=8)
            if header:
                pdf.set_xy(36, 28)
                pdf.cell(text=header)
            pdf.set_y(60)
            for line in body:
                text, size = line if isinstance(line, tuple) else (line, 8)
                pdf.set_font("Helvetica", size=size)
                if text:
                    pdf.cell(h=size + 1, text=text, new_x="LMARGIN", new_y="NEXT")
                else:
                    pdf.ln(9)
        pdf.output(str(path))
        return path

    return make


@pytest.fixture(scope="session")
def timetable_block() -> Callable:
    # Lays out a block at the top of a sheet as the 2026 timetable workbook does: its title across
    # the grid, each semester's heading and days, and the handbook's period key (p. 10) down each
    # semester's side, with periods 1 to 5, the meridian and 6 to 9 from row 5
    key = (("1", 8), ("2", 9), ("3", 10), ("4", 11), ("5", 12), ("M", 13), ("6", 14))
    key += (("7", 15), ("8", 16), ("9", 17))
    days = ("MONDAY", "TUESDAY", "WEDNESDAY", "THURSDAY", "FRIDAY")

    def block(sheet, title: str) -> None:
        sheet["A1"] = title
        sheet.merge_cells("A1:O1")
        for column, semester in ((1, "Semester 1"), (9, "Semester 2")):
            sheet.cell(2, column, f"Mainstream EE - {semester}")
            sheet.cell(3, column, semester)
            for offset, day in enumerate(days):
                sheet.cell(3, column + 2 + offset, day)
            sheet.cell(4, column, "Start Time")
            sheet.cell(4, column + 1, "Period")
            for row, (period, hour) in enumerate(key, start=5):
                sheet.cell(row, column, time(hour))
                sheet.cell(row, column + 1, period if period == "M" else int(period))

    return block


@pytest.fixture(scope="session")
def complementary_xlsx() -> Callable:
    # Builds a complementary workbook in the real one's layout: each sheet's title, its header, then
    # its rows, with one cell changed afterwards where a test plants a fault
    def make(path: Path, rows: dict[str, list[tuple]], change=None) -> Path:
        book = openpyxl.Workbook()
        book.remove(book.active)
        for name, sheet_rows in rows.items():
            sheet = book.create_sheet(name)
            sheet.append([f"The {name} sheet's title"])
            sheet.append(
                list(MEETING_COLUMNS if name == "Meetings" else OFFERING_COLUMNS)
            )
            for row in sheet_rows:
                sheet.append(list(row))
        if change:
            name, cell, value = change
            book[name][cell] = value
        book.save(path)
        return path

    return make


@pytest.fixture(scope="session")
def pdf_of() -> Callable[[list[list[str]]], bytes]:
    # Prints pages of lines into a PDF, one line per row. Windows-1252 lets the built-in font print
    # the record's dashes.
    def make(pages: list[list[str]]) -> bytes:
        pdf = FPDF()
        pdf.core_fonts_encoding = "windows-1252"
        pdf.set_font("Helvetica", size=9)
        for page in pages:
            pdf.add_page()
            for line in page:
                pdf.cell(text=line, new_x="LMARGIN", new_y="NEXT")
        return bytes(pdf.output())

    return make
