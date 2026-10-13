# Checks ATP-1's label sheets before each is committed: every cell I fill holds a value of the right
# kind, and every tree is one EntryRule accepts, in the shape E2's prompt.md asks of every label
# (experiments/ingestion/requirement-parsing/). Trees are read with EntryRule alone, never SAGE's
# grammar or readers, so nothing SAGE produces stands between a printed page and its label. It also
# holds each sheet's layout, which make_sheets.py builds from and the harness will read by.
# From the repo root:
#   uv run --project backend python tests/acceptance/atp1/labels/check_labels.py <sheet.xlsx> ...
import argparse
import json
import re
import sys
from collections.abc import Callable, Iterator
from datetime import time
from pathlib import Path
from typing import Any, NamedTuple, get_args

from openpyxl import load_workbook
from pydantic import TypeAdapter, ValidationError

from sage.domain.rules import EntryRule, TermCode

Row = dict[str, Any]

CODE = re.compile(r"[A-Z]{3}\d{4}[A-Z]")
DAYS = ("Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday")
# The columns whose value comes from a fixed list, which make_sheets.py offers as a dropdown. The
# NQF levels and kinds are the rule-tables migration's checks on courses and timetable_slots.
CHOICES = {
    "NQF level": tuple(str(level) for level in range(5, 11)),
    "Term": get_args(TermCode),
    "Day": DAYS,
    "Kind": (
        "lecture",
        "tutorial",
        "practical",
        "lab",
        "workshop",
        "screening",
        "conversation",
    ),
    "Ambiguous": ("yes",),
}
# What make_sheets.py writes in Note for a course a curriculum lists but the book never describes
NO_ENTRY = "no entry printed"
ENTRY_RULE = TypeAdapter(EntryRule | None)


def required(row: Row, *columns: str) -> Iterator[str]:
    for column in columns:
        if row[column] is None:
            yield f"{column} is empty"


def together(row: Row, *columns: str) -> Iterator[str]:
    filled = [row[column] is not None for column in columns]
    if any(filled) and not all(filled):
        yield f"{', '.join(columns[:-1])} and {columns[-1]} go together: fill all or none"


def whole_number(row: Row, column: str) -> Iterator[str]:
    if row[column] is not None and not str(row[column]).isdigit():
        yield f"{column} isn't a whole number"


def choice(row: Row, column: str) -> Iterator[str]:
    if row[column] is not None and str(row[column]) not in CHOICES[column]:
        yield f"{column} isn't one of {', '.join(CHOICES[column])}"


def course_codes(row: Row) -> Iterator[str]:
    codes = row["Course"] and str(row["Course"]).split(",")
    if codes and not all(CODE.fullmatch(code.strip()) for code in codes):
        yield "Course isn't full course codes separated by commas"


def days(row: Row) -> Iterator[str]:
    names = row["Days"] and str(row["Days"]).split(",")
    if names and not all(name.strip() in DAYS for name in names):
        yield "Days isn't weekday names separated by commas"


def clock(value: Any) -> time:
    # A time typed as HH:MM into a text cell, or one Excel has turned into a time
    return value if isinstance(value, time) else time.fromisoformat(str(value))


def times(row: Row) -> Iterator[str]:
    if row["Start"] is None or row["End"] is None:
        return
    try:
        start, end = clock(row["Start"]), clock(row["End"])
    except ValueError:
        yield "Start and End must be times such as 08:00"
        return
    if end <= start:
        yield "End isn't after Start"


def node_problems(node: dict[str, Any]) -> Iterator[str]:
    # What EntryRule lets through but would later count against SAGE: a course leaf whose code
    # isn't whole, such as EEE3089, and the shapes prompt.md's rule 10 forbids
    if "course" in node and not CODE.fullmatch(node["course"]):
        yield f"names {node['course']}, which isn't a full course code"
    for kind in ("all", "any"):
        children = node.get(kind, [])
        if kind in node and len(children) < 2:
            yield f"has an {kind} with one child"
        if any(kind in child for child in children):
            yield f"has an {kind} directly inside an {kind}"
        for child in children:
            yield from node_problems(child)


def tree_problems(tree: Any) -> Iterator[str]:
    # null is the rule of a field that holds no requirement, as in prompt.md
    try:
        rule = ENTRY_RULE.validate_json(str(tree))
    except ValidationError as error:
        found = (
            f"{'.'.join(map(str, problem['loc'])) or 'tree'}: {problem['msg']}"
            for problem in error.errors()
        )
        yield f"isn't one EntryRule accepts ({'; '.join(found)})"
        return
    if rule is not None:
        yield from node_problems(json.loads(str(tree)))


def requirement(row: Row, column: str, required: bool) -> Iterator[str]:
    # A printed requirement and its tree, which is needed unless nothing is printed
    printed, tree = row[column], row[f"{column} tree"]
    if printed is None and required:
        yield f"{column} is empty: write None when nothing is printed"
    if printed is None or str(printed).rstrip(".").lower() == "none":
        if tree is not None and str(tree) != "null":
            yield f"{column} tree must be blank or null when nothing is printed"
    elif tree is None:
        yield f"{column} tree is empty: write null if the print holds no requirement"
    else:
        yield from (f"{column} tree {problem}" for problem in tree_problems(tree))


def course_problems(row: Row) -> Iterator[str]:
    if row["Page"] is None:
        # The book never describes the course, so the harness expects SAGE to leave it out
        others = [row[column] for column in COURSES.mine if column != "Note"]
        if row["Note"] != NO_ENTRY or any(value is not None for value in others):
            yield f'a course with no page holds only "{NO_ENTRY}" in Note'
        return
    yield from required(row, "Title", "Credits", "NQF level", "Term")
    yield from whole_number(row, "Credits")
    yield from choice(row, "NQF level")
    yield from choice(row, "Term")
    yield from requirement(row, "Requirement", required=True)
    yield from requirement(row, "Co-requisite", required=False)


def timetable_problems(row: Row) -> Iterator[str]:
    # A cell that names a course but isn't a meeting, such as a note, has no day, times or term
    yield from required(row, "Course")
    yield from course_codes(row)
    yield from together(row, "Day", "Start", "End", "Term")
    yield from choice(row, "Day")
    yield from times(row)
    yield from choice(row, "Kind")
    yield from choice(row, "Term")


def complementary_problems(row: Row) -> Iterator[str]:
    yield from required(row, "Book", "Page")
    yield from whole_number(row, "Page")
    yield from days(row)
    yield from together(row, "Start", "End")
    yield from times(row)
    yield from choice(row, "Kind")
    yield from choice(row, "Ambiguous")


def held_out_problems(row: Row) -> Iterator[str]:
    yield from requirement(row, "Requirement", required=True)
    yield from requirement(row, "Co-requisite", required=False)


class Sheet(NamedTuple):
    given: tuple[str, ...]  # what names each row, which make_sheets.py writes
    mine: tuple[str, ...]  # what I label from print
    problems: Callable[[Row], Iterator[str]]


COURSES = Sheet(
    ("Code", "Page"),
    (
        "Title",
        "Credits",
        "NQF level",
        "Term",
        "Requirement",
        "Requirement tree",
        "Co-requisite",
        "Co-requisite tree",
        "Note",
    ),
    course_problems,
)
SHEETS = {
    "courses_2026.xlsx": COURSES,
    "courses_2023.xlsx": COURSES,
    "timetable.xlsx": Sheet(
        ("Sheet", "Cell"),
        ("Course", "Section", "Day", "Start", "End", "Kind", "Term", "Note"),
        timetable_problems,
    ),
    "complementary.xlsx": Sheet(
        ("Row", "Course", "Place"),
        ("Days", "Start", "End", "Kind", "Book", "Page", "Ambiguous", "Note"),
        complementary_problems,
    ),
    "held_out.xlsx": Sheet(
        ("Code", "Page"),
        (
            "Requirement",
            "Requirement tree",
            "Co-requisite",
            "Co-requisite tree",
            "Note",
        ),
        held_out_problems,
    ),
}


def blank_as_none(value: Any) -> Any:
    # Excel keeps spaces typed into a cell, so a cell of spaces is blank too
    if isinstance(value, str):
        return value.strip() or None
    return value


def read(path: Path) -> dict[int, Row]:
    # A sheet's rows by their row number, keyed by column, leaving out rows with nothing in them.
    # A sheet whose columns differ from its layout is refused, so no label lands in the wrong place.
    sheet = SHEETS[path.name]
    columns = sheet.given + sheet.mine
    cells = load_workbook(path)["Labels"].iter_rows(values_only=True)
    if next(cells) != columns:
        raise ValueError(f"{path.name}'s columns aren't {', '.join(columns)}")
    rows = {
        number: dict(zip(columns, map(blank_as_none, values), strict=True))
        for number, values in enumerate(cells, start=2)
    }
    return {
        number: row
        for number, row in rows.items()
        if any(value is not None for value in row.values())
    }


def problems_in(sheet: Sheet, row: Row) -> list[str]:
    if all(row[column] is None for column in sheet.mine):
        return ["not labelled"]
    return list(sheet.problems(row))


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Names every row of ATP-1's label sheets whose labels are missing "
        "or malformed, and fails if there is one."
    )
    parser.add_argument("sheets", nargs="+", type=Path, help="the .xlsx files to check")
    failed = False
    for path in parser.parse_args().sheets:
        sheet = SHEETS[path.name]
        rows = read(path)
        found = {number: problems_in(sheet, row) for number, row in rows.items()}
        failing = {number: problems for number, problems in found.items() if problems}
        summary = f"{len(failing)} with problems" if failing else "all pass"
        print(f"{path.name}: {len(rows)} rows, {summary}")
        for number, problems in failing.items():
            row = rows[number]
            name = " ".join(
                str(row[column]) for column in sheet.given[:2] if row[column]
            )
            print(f"  row {number} ({name}): {'; '.join(problems)}")
        failed = failed or bool(failing)
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
