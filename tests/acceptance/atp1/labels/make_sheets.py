# Makes ATP-1's blank label sheets from the sources alone. Each sheet gets only what names its rows
# (a course code and the page of its entry, a timetable cell, a workbook row), never a value I
# label, and nothing here uses SAGE's readers, grammar or output. It prints nothing, so no page of
# a held-out department is ever shown, and it refuses to write over a sheet, so a rerun can't wipe
# labels. Run it into an empty folder to check a draw, from the repo root:
#   uv run --project backend python tests/acceptance/atp1/labels/make_sheets.py <folder>
import argparse
import hashlib
import random
import re
from pathlib import Path
from typing import Any

import pdfplumber
from check_labels import CHOICES, NO_ENTRY, SHEETS, Row
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[4]
SOURCES = ROOT / "data" / "sources"
# Printed pages from each book's contents: the six Electrical Engineering curricula, then every
# course entry, from the faculty's first department to the last of the other faculties
BOOKS = {2026: ((48, 61), (71, 195)), 2023: ((46, 59), (69, 197))}
# The timetable sheets the 2026 manifest reads
TIMETABLE = ("EB009", "EB011", "EB022", "EB809", "EB811", "EB822", "Labs")
# The book's code key gives these prefixes to Chemical Engineering and Construction Economics
HELD_OUT = ("CHE", "CON")
# Seen before the two departments were held out, as docs/shortfalls.md records
SEEN = {"CON2026S", "CON3053S", "CON3054S"}
# Any number fixed before the draws would do; this one is for EEE4022S
SEED = 4022

CODE = re.compile(r"\b[A-Z]{3}\d{4}[A-Z]\b")
# Such as CSC3xxx, any third-year CSC course
PLACEHOLDER = re.compile(r"\b([A-Z]{3}\d)xxx\b")
# Such as EEE4022F/S, which heads both terms' entry
HEADING = re.compile(r"[A-Z]{3}\d{4}[A-Z](/[A-Z])?\b")
NAMES_A_COURSE = re.compile(r"[A-Z]{3}\d{4}")  # also EEE3097/8S, which names two
ENTRY_RULE = (
    "The printed page of the entry's heading. Each entry prints one 'NQF credits at NQF level' "
    "line, and its heading is the last bold line before it that starts with a course code, or "
    "the last such line if none is bold. A heading such as EEE4022F/S names both courses."
)
WIDTHS = {"Title": 40, "Requirement": 60, "Co-requisite": 40, "Book": 40, "Note": 40}


def fingerprint(path: Path) -> str:
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    return f"{path.relative_to(ROOT)}, SHA-256 {digest}"


def page_lines(pdf: pdfplumber.PDF, number: int) -> list[dict[str, Any]]:
    # Printed page n is the PDF's page n in both books, as each page's running header shows
    lines = pdf.pages[number].extract_text_lines(return_chars=True)
    if str(number) not in re.findall(r"\d+", lines[0]["text"]):
        raise ValueError(f"PDF page {number} isn't printed page {number}")
    return lines


def listed(pdf: pdfplumber.PDF, first: int, last: int) -> tuple[set[str], set[str]]:
    # Every course code the six curricula print, in their tables and notes, from the first one's
    # heading to the transferees' section after the last, and their placeholders' stems
    lines = [
        line["text"] for n in range(first, last + 1) for line in page_lines(pdf, n)
    ]
    start = next(i for i, line in enumerate(lines) if "[EB009EEE01]" in line)
    end = next(
        i
        for i, line in enumerate(lines)
        if line.startswith("Access for University of Technology Transferees")
    )
    curricula = "\n".join(lines[start:end])
    return set(CODE.findall(curricula)), set(PLACEHOLDER.findall(curricula))


def entries(pdf: pdfplumber.PDF, first: int, last: int) -> dict[str, int]:
    # Each entry's heading as printed, with its page. A bold heading wins over a later line that
    # starts with a code; the fallback catches HUB4049F's 2026 heading, which isn't bold.
    found: dict[str, int] = {}
    bold = plain = None
    for number in range(first, last + 1):
        for line in page_lines(pdf, number):
            if heading := HEADING.match(line["text"]):
                plain = (heading.group(), number)
                if "Bold" in line["chars"][0]["fontname"]:
                    bold = plain
            if "NQF credits at NQF level" in line["text"]:
                if plain is None:
                    raise ValueError(f"page {number}: a credits line with no heading")
                heading_text, page = bold or plain
                found.setdefault(heading_text, page)
                bold = plain = None
    return found


def courses_named(heading: str) -> list[str]:
    # EEE4022F/S names EEE4022F and EEE4022S
    return [heading[:7] + suffix for suffix in heading[7:].split("/")]


def draw[T](population: list[T], size: int) -> list[T]:
    # The same seed gives the same draw on every run of Python 3.12, kept in population order
    drawn = set(random.Random(SEED).sample(population, size))
    return [item for item in population if item in drawn]


def write(
    folder: Path, name: str, rows: list[Row], about: list[tuple[str, str]]
) -> None:
    sheet = SHEETS[name]
    columns = sheet.given + sheet.mine
    workbook = Workbook()
    labels = workbook.active
    labels.title = "Labels"
    labels.append(columns)
    for row in rows:
        labels.append([row.get(column) for column in columns])
    for cell in labels[1]:
        cell.font = Font(bold=True)
    # The header and what names each row stay in view while scrolling
    labels.freeze_panes = labels.cell(2, len(sheet.given) + 1)
    for index, column in enumerate(columns, start=1):
        letter = get_column_letter(index)
        labels.column_dimensions[letter].width = WIDTHS.get(
            column.removesuffix(" tree"), 14
        )
        if column not in sheet.mine:
            continue
        # Text, so Excel keeps a time or a code as typed rather than converting it
        for (cell,) in labels.iter_rows(min_row=2, min_col=index, max_col=index):
            cell.number_format = "@"
        if column in CHOICES:
            dropdown = DataValidation(
                type="list", formula1=f'"{",".join(CHOICES[column])}"', allow_blank=True
            )
            dropdown.add(f"{letter}2:{letter}{len(rows) + 1}")
            labels.add_data_validation(dropdown)
    notes = workbook.create_sheet("About")
    for line in about:
        notes.append(line)
    notes.column_dimensions["A"].width = 12
    notes.column_dimensions["B"].width = 120
    with (folder / name).open("xb") as file:
        workbook.save(file)


def make_courses(
    folder: Path, year: int, path: Path, pdf: pdfplumber.PDF
) -> dict[str, int]:
    (first, last), entry_pages = BOOKS[year]
    codes, stems = listed(pdf, first, last)
    headings = entries(pdf, *entry_pages)
    pages: dict[str, int] = {}
    for heading, page in headings.items():
        for code in courses_named(heading):
            pages.setdefault(code, page)
    codes |= {code for code in pages if code[:4] in stems}
    rows = [
        {"Code": code, "Page": pages[code]}
        if code in pages
        else {"Code": code, "Note": NO_ENTRY}
        for code in codes
    ]
    rows.sort(
        key=lambda row: (row.get("Page") is None, row.get("Page", 0), row["Code"])
    )
    placeholders = ", ".join(f"{stem}xxx" for stem in sorted(stems))
    write(
        folder,
        f"courses_{year}.xlsx",
        rows,
        [
            ("Source", fingerprint(path)),
            (
                "Rows",
                (
                    f"Every course the six Electrical Engineering curricula print in their "
                    f"tables or notes, on printed pages {first} to {last}, with {placeholders} "
                    f"read as each entry in the book it matches: {len(rows)} courses, sorted "
                    "by page"
                ),
            ),
            ("Page", ENTRY_RULE),
            (
                "No page",
                (
                    f'The book prints no entry for the course, so Note says "{NO_ENTRY}" and '
                    "the other cells stay blank"
                ),
            ),
        ],
    )
    return headings


def make_held_out(folder: Path, path: Path, headings: dict[str, int]) -> None:
    rows = [
        {"Code": heading, "Page": page}
        for heading, page in headings.items()
        if heading.startswith(HELD_OUT)
        and not SEEN.intersection(courses_named(heading))
    ]
    rows.sort(key=lambda row: (row["Page"], row["Code"]))
    write(
        folder,
        "held_out.xlsx",
        rows,
        [
            ("Source", fingerprint(path)),
            (
                "Rows",
                (
                    "Every entry whose code starts CHE (Chemical Engineering) or CON "
                    f"(Construction Economics and Management), less {', '.join(sorted(SEEN))}, "
                    f"seen before the two departments were held out: {len(rows)} entries, "
                    "sorted by page. All are labelled, so there is no draw"
                ),
            ),
            ("Page", ENTRY_RULE),
        ],
    )


def make_timetable(folder: Path) -> None:
    path = SOURCES / "2026" / "2026_timetable.xlsx"
    workbook = load_workbook(path)
    # A merged block's value sits in its first cell only, so each block counts once
    cells = [
        (name, cell.coordinate)
        for name in TIMETABLE
        for row in workbook[name].iter_rows()
        for cell in row
        if cell.value is not None and NAMES_A_COURSE.search(str(cell.value))
    ]
    rows = [{"Sheet": name, "Cell": coordinate} for name, coordinate in draw(cells, 60)]
    write(
        folder,
        "timetable.xlsx",
        rows,
        [
            ("Source", fingerprint(path)),
            (
                "Population",
                (
                    f"Every cell holding a course code on {', '.join(TIMETABLE)}, a merged "
                    f"block counting once: {len(cells)} cells, in sheet order, then row by row"
                ),
            ),
            (
                "Draw",
                (
                    f"random.Random({SEED}).sample(population, 60) on Python 3.12, "
                    "kept in population order"
                ),
            ),
        ],
    )


def make_complementary(folder: Path) -> None:
    path = SOURCES / "2026" / "2026_Complementary_Studies_Course_Details.xlsx"
    meetings = load_workbook(path)["Meetings"]
    # Row 1 holds the sheet's title and row 2 its column names
    codes = {
        number: code
        for number, (code,) in enumerate(
            meetings.iter_rows(min_row=3, max_col=1, values_only=True), start=3
        )
    }
    rows = []
    for number in draw(list(codes), 40):
        same = [n for n, code in codes.items() if code == codes[number]]
        place = f"{same.index(number) + 1} of {len(same)}"
        rows.append({"Row": number, "Course": codes[number], "Place": place})
    write(
        folder,
        "complementary.xlsx",
        rows,
        [
            ("Source", fingerprint(path)),
            (
                "Population",
                f"Rows 3 to {meetings.max_row} of the Meetings sheet: {len(codes)} rows",
            ),
            (
                "Draw",
                (
                    f"random.Random({SEED}).sample(population, 40) on Python 3.12, "
                    "kept in row order"
                ),
            ),
            (
                "Place",
                (
                    "The row's place among its course's rows, in the workbook's order: 2 of 3 "
                    "is the course's second row of three"
                ),
            ),
        ],
    )


def main() -> None:
    parser = argparse.ArgumentParser(description="Makes ATP-1's blank label sheets.")
    parser.add_argument("folder", type=Path, help="where to write the sheets")
    folder = parser.parse_args().folder
    for year in BOOKS:
        path = SOURCES / str(year) / f"EBE_{year}.pdf"
        with pdfplumber.open(path) as pdf:
            headings = make_courses(folder, year, path, pdf)
        if year == 2026:
            make_held_out(folder, path, headings)
    make_timetable(folder)
    make_complementary(folder)


if __name__ == "__main__":
    main()
