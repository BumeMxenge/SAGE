# Draws ATP-3's human sample and reads my marks back. The first run draws 40 judged replies, shares
# the places between the (arm, judge outcome) strata with measures.allocation, shuffles the draw with
# a fixed seed, and writes results/human_sample.xlsx for me to mark. The sheet shows each reply with
# what a right one does, but never a mark of the judge's, so I mark blind to it. No run writes over
# the sheet: once it exists, a run reads my marks into results/human.csv instead, for report.py to
# set beside the judge's. make atp3-report runs it from backend/, before report.py, and calls no
# model.
import csv
import json
import random
import re
from collections import defaultdict
from collections.abc import Sequence
from pathlib import Path
from typing import Any, get_args

from ask import SEED, as_json, write
from judge import Outcome
from measures import allocation
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from resolve import RESULTS, Row, read_prompts, read_resolved

SIZE = 40  # replies I mark by hand
OUTCOMES = get_args(Outcome)
YES_NO = ("yes", "no")
# What each row shows before my marks: the question, what a right reply does and the reply itself.
# Nothing the judge decided, as its marks would sway mine.
SHOWN = (
    "Sample",
    "Turn",
    "Arm",
    "Question",
    "Expected Behaviour",
    "Expected Fact",
    "Handbook Quote",
    "Page",
    "Trap",
    "Engine and Plan Facts",
    "Reply",
    "Referral Flag",
    "Plan Offered",
    "Citations",
)
SUPPORTS = re.compile(r"Citation (\d+) Supports")
# In characters: long text gets room, and the rest is wide enough for its values
WIDTHS = {
    "Turn": 18,
    "Question": 40,
    "Expected Fact": 40,
    "Handbook Quote": 40,
    "Trap": 40,
    "Engine and Plan Facts": 50,
    "Reply": 70,
    "Citations": 70,
    "My Outcome": 18,
    "My Notes": 40,
}
COLUMNS = ("sample", "turn", "arm", "my_outcome", "my_marks", "notes")


def read_results(name: str) -> list[Row]:
    with (RESULTS / f"{name}.csv").open(encoding="utf-8", newline="") as file:
        return list(csv.DictReader(file))


def draw(judged: Sequence[Row]) -> list[str]:
    # The turns to mark. Each stratum is drawn in sorted order from one generator, so the same
    # results always give the same sample. The shuffle after it keeps the sheet from grouping
    # replies by the judge's outcome, which I mark blind to.
    strata: dict[str, list[str]] = defaultdict(list)
    for row in judged:
        strata[f"{row['arm']}/{row['judge_outcome']}"].append(row["turn"])
    places = allocation({name: len(turns) for name, turns in strata.items()}, SIZE)
    rng = random.Random(SEED)
    sample = [
        turn
        for name in sorted(strata)
        for turn in rng.sample(sorted(strata[name]), places[name])
    ]
    rng.shuffle(sample)
    return sample


def dropdown(options: Sequence[str]) -> DataValidation:
    # showErrorMessage makes Excel refuse a typo, where it would otherwise only offer the list
    return DataValidation(
        type="list",
        formula1=f'"{",".join(options)}"',
        allow_blank=True,
        showErrorMessage=True,
    )


def write_sheet(path: Path, sample: Sequence[str]) -> None:
    turns = {row["turn"]: row for row in read_results("turns")}
    prompts = {row["id"]: row for row in read_prompts()}
    resolved = read_resolved()
    cited = {name: json.loads(turns[name]["citations"] or "[]") for name in sample}
    # A Supports column for each citation of the reply with the most, and at least one
    widest = max([1, *map(len, cited.values())])
    columns = (
        *SHOWN,
        "My Outcome",
        *(f"Citation {n} Supports" for n in range(1, widest + 1)),
        "My Notes",
    )
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Replies"
    sheet.append(columns)
    for number, name in enumerate(sample, start=1):
        turn = turns[name]
        prompt, fixed = prompts[turn["prompt"]], resolved[turn["prompt"]]
        hidden = " (hidden)" if turn["offered_hidden"] == "yes" else ""
        cells = {
            "Sample": number,
            "Turn": name,
            "Arm": turn["arm"],
            "Question": prompt["prompt"],
            "Expected Behaviour": fixed["expected"],
            "Expected Fact": prompt["expected_fact"],
            "Handbook Quote": prompt["quote"],
            "Page": prompt["page"],
            "Trap": prompt["tempting_wrong"],
            "Engine and Plan Facts": fixed["words"],
            "Reply": turn["reply"],
            "Referral Flag": turn["referral"],
            "Plan Offered": turn["offered"] + hidden,
            "Citations": "\n\n".join(
                f'{c["position"]}. "{c["quote"]}" ({c["section"]}, page {c["page"]})\n'
                f"Passage: {c['text']}"
                for c in cited[name]
            ),
        }
        # My columns are left out of cells, so they stay blank
        sheet.append([cells.get(column) for column in columns])
    outcome, yes_no = dropdown(OUTCOMES), dropdown(YES_NO)
    for index, column in enumerate(columns, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = WIDTHS.get(column, 14)
        if column == "My Outcome":
            outcome.add(f"{letter}2:{letter}{len(sample) + 1}")
        elif SUPPORTS.fullmatch(column):
            yes_no.add(f"{letter}2:{letter}{len(sample) + 1}")
    sheet.add_data_validation(outcome)
    sheet.add_data_validation(yes_no)
    for line in sheet.iter_rows():
        for cell in line:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    # The header, and which reply a row is, stay in view while I scroll to my columns on the right
    sheet.freeze_panes = "D2"
    # "xb" refuses a file that exists, so even a slip in main can't wipe my marks
    with path.open("xb") as file:
        workbook.save(file)
    print(f"results/{path.name}: {len(sample)} replies. Mark them, then run it again.")


def text(cell: Any) -> str:
    # A cell as I left it: empty is None, and Excel keeps any spaces I typed
    return "" if cell is None else str(cell).strip()


def chosen(sample: Any, column: str, cell: Any, options: Sequence[str]) -> str:
    # What I picked from a dropdown's list. Pasting skips Excel's check, so this is the real one.
    value = text(cell)
    if value and value not in options:
        raise SystemExit(
            f"Sample {sample}: {column} is {value!r}, not one of {', '.join(options)}"
        )
    return value


def read_marks(path: Path) -> list[Row]:
    # My marks, each column found by its header text so one I've moved still reads right
    lines = load_workbook(path)["Replies"].iter_rows(values_only=True)
    column = {str(name): index for index, name in enumerate(next(lines))}
    supports = sorted(
        (int(found[1]), index)
        for name, index in column.items()
        if (found := SUPPORTS.fullmatch(name))
    )
    rows = []
    for line in lines:
        # Formatting Excel keeps below the last reply leaves rows with nothing in them
        if not any(line):
            continue
        sample = line[column["Sample"]]
        outcome = chosen(sample, "My Outcome", line[column["My Outcome"]], OUTCOMES)
        marks = []
        for position, index in supports:
            mark = chosen(sample, f"Citation {position} Supports", line[index], YES_NO)
            if mark:
                marks.append({"position": position, "supports": mark == "yes"})
        rows.append(
            {
                "sample": str(sample),
                "turn": text(line[column["Turn"]]),
                "arm": text(line[column["Arm"]]),
                "my_outcome": outcome,
                "my_marks": as_json(marks),
                "notes": text(line[column["My Notes"]]),
            }
        )
    return rows


def main() -> None:
    path = RESULTS / "human_sample.xlsx"
    if not path.exists():
        # A turn with no judge outcome never reached her, so there's no reply to mark
        judged = [row for row in read_results("judge") if row["judge_outcome"]]
        write_sheet(path, draw(judged))
        return
    rows = read_marks(path)
    write("human", COLUMNS, rows)
    filled = sum(bool(row["my_outcome"]) for row in rows)
    print(f"{filled} of {len(rows)} rows have My Outcome filled")


if __name__ == "__main__":
    main()
